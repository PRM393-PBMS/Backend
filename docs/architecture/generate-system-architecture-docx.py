from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "system-architecture.docx"
DIAGRAM = HERE / "system-architecture.html"

NAVY = "15324B"
BLUE = "1976A5"
TEAL = "168B75"
PALE_BLUE = "EAF3F8"
PALE_GRAY = "F3F5F7"
PALE_AMBER = "FFF4D8"
DARK = "243746"
MUTED = "5E6C78"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=95, start=110, bottom=95, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell_text(cell, text, *, bold=False, color=DARK, size=9):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, title in enumerate(headers):
        set_cell_text(header.cells[i], title, bold=True, color="FFFFFF", size=9)
        set_cell_shading(header.cells[i], NAVY)
        if widths:
            header.cells[i].width = Inches(widths[i])
    for row_num, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=font_size)
            if row_num % 2 == 1:
                set_cell_shading(cells[i], PALE_GRAY)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_callout(doc, title, body, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    r.font.size = Pt(10.5)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    for run in p2.runs:
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_hyperlink(paragraph, text, target):
    part = paragraph.part
    rel_id = part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (("Title", 30, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 10.5, TEAL)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name in ("Title", "Heading 1", "Heading 2") else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(11 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "PRM393  /  PARKING MANAGEMENT                                      SYSTEM ARCHITECTURE"
    header.style = doc.styles["Caption"]
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(section.footer.paragraphs[0])


def build_document():
    doc = Document()
    style_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(76)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("PRM393 · PARKING MANAGEMENT")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(9)
    title.add_run("Tài liệu kiến trúc\nhệ thống")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    sr = subtitle.add_run("Mô tả đầy đủ System Architecture Diagram")
    sr.font.size = Pt(15)
    sr.font.color.rgb = RGBColor.from_string(MUTED)

    add_callout(
        doc,
        "Phạm vi tài liệu",
        "Ứng dụng Flutter Android/iOS, backend NestJS + Express chạy trên Render, các domain module, dữ liệu Supabase PostgreSQL, tích hợp PayOS/Resend/Plate Recognizer, bảo mật, batch job, monitoring và các giới hạn đã xác nhận trong codebase.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.add_run("Phiên bản 1.0  ·  ").bold = True
    p.add_run(date.today().strftime("%d/%m/%Y"))
    p = doc.add_paragraph()
    p.add_run("Sơ đồ tương tác: ").bold = True
    add_hyperlink(p, "system-architecture.html", "system-architecture.html")
    p = doc.add_paragraph()
    p.add_run("Nguồn sơ đồ: ").bold = True
    p.add_run("system-architecture.json")

    doc.add_page_break()

    # Contents and summary
    doc.add_heading("Mục lục", level=1)
    for item in [
        "1. Tổng quan và phạm vi",
        "2. Ranh giới hệ thống và triển khai",
        "3. Thành phần kiến trúc",
        "4. Luồng dữ liệu và giao thức",
        "5. Lưu trữ dữ liệu",
        "6. Bảo mật",
        "7. Vận hành, monitoring và giới hạn hạ tầng",
        "8. Các thành phần không có trong runtime",
        "9. Bằng chứng trong codebase",
        "10. Thuật ngữ",
    ]:
        p = doc.add_paragraph(item)
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(2)

    doc.add_heading("1. Tổng quan và phạm vi", level=1)
    doc.add_paragraph(
        "Hệ thống quản lý bãi đỗ xe có kiến trúc client–server. Client trong phạm vi dự án là ứng dụng Flutter cho Android/iOS. Ứng dụng gửi JSON qua HTTPS bằng Dio đến REST API NestJS + Express. Backend chạy như một production Web Service trên Render; các domain module cùng nằm trong một process. PostgreSQL trên Supabase là nơi lưu dữ liệu nghiệp vụ lâu dài, truy cập qua Prisma và Session Pooler."
    )
    doc.add_paragraph(
        "Các dịch vụ ngoài được backend gọi trực tiếp qua HTTPS gồm PayOS (thanh toán), Resend (email OTP) và Plate Recognizer (OCR biển số). UptimeRobot kiểm tra endpoint health từ bên ngoài. Trong codebase không có queue, worker riêng, cache, search engine, object storage hay API gateway độc lập."
    )
    add_callout(
        doc,
        "Cách đọc các ranh giới",
        "Render — production Web Service là ranh giới triển khai bên ngoài. Application trust boundary nằm bên trong Render và biểu thị vùng runtime ứng dụng cùng cấu hình secrets. Domain services là nhóm module nghiệp vụ bên trong backend; Build / deployment thuộc Render nhưng nằm ngoài trust boundary ứng dụng.",
        PALE_AMBER,
    )

    doc.add_heading("2. Ranh giới hệ thống và triển khai", level=1)
    add_table(doc, ["Vùng", "Thành phần nằm trong vùng", "Ý nghĩa"], [
        ("Flutter mobile application", "Flutter UI, Dio API client, secure token storage", "Ứng dụng Android/iOS và lưu token trên thiết bị; client không nằm trong repository backend."),
        ("Render — production Web Service", "NestJS runtime, các module, cron, local state, build/deployment", "Một dịch vụ/process production; không có worker hoặc gateway riêng được cấu hình."),
        ("Application trust boundary", "Auth, API, domain modules, Prisma, cron và volatile state", "Vùng ứng dụng dùng cấu hình môi trường; secrets không được đưa vào sơ đồ."),
        ("Domain services", "Identity/profile, catalog/parking, booking, billing, operations", "Các nhóm nghiệp vụ được gọi bởi REST controllers và dùng lớp truy cập dữ liệu."),
        ("Supabase managed data platform", "Session Pooler và PostgreSQL", "Dịch vụ dữ liệu được quản lý; PostgreSQL là system of record."),
        ("External integrations", "PayOS, Resend, Plate Recognizer", "Các API bên thứ ba được backend gọi qua HTTPS."),
        ("External monitoring", "UptimeRobot", "Kiểm tra HTTP liveness từ ngoài hệ thống."),
    ], widths=[1.65, 2.55, 2.65])

    doc.add_heading("3. Thành phần kiến trúc", level=1)
    add_table(doc, ["Thành phần", "Vai trò và nội dung chính", "Vị trí / lưu ý"], [
        ("Flutter UI", "Màn hình ứng dụng Dart cho Android và iOS.", "Client mobile duy nhất được xác định trong phạm vi này."),
        ("Dio API client", "Gửi/nhận JSON REST qua HTTPS; timeout và interceptors; gắn Bearer token theo cấu hình client.", "Trong ứng dụng Flutter."),
        ("Secure token storage", "Lưu access/refresh token trên thiết bị.", "Không phải kho secrets của backend."),
        ("NestJS REST controllers", "HTTP entry `/api/*`; controllers, DTO validation và Swagger `/api/docs`.", "Chạy trong Render Web Service."),
        ("Auth module + guards", "JWT access/refresh, global AtGuard và role checks.", "Route public là ngoại lệ được khai báo; không có auth provider SaaS riêng trong sơ đồ."),
        ("Identity & profile", "Các module Auth, Users và Roles.", "OTP được quản lý bởi dịch vụ OTP trong process và gửi bằng Resend."),
        ("Catalog & parking", "Catalog, parking operations, QR và luồng OCR biển số.", "Upload ảnh có thể đi qua filesystem local trước khi gọi OCR."),
        ("Booking & subscriptions", "Reservations và subscriptions.", "Reservation có tác vụ cập nhật trạng thái chạy theo lịch."),
        ("Billing services", "Payments, wallets và renewals.", "Tích hợp PayOS qua API HTTPS và cập nhật DB."),
        ("Operations services", "Incidents, reports và PDF.", "Một số luồng upload proof dùng filesystem local."),
        ("Prisma data access", "PrismaModule dùng Prisma Client để đọc/ghi dữ liệu.", "Kết nối PostgreSQL qua Supabase Session Pooler/TLS."),
        ("Reservation cron", "Nest Schedule chạy mỗi phút trong process backend.", "Không phải worker/service độc lập."),
        ("Volatile local state", "OTP Map TTL 10 phút và filesystem `UPLOAD_DIR`.", "Không phải Redis hay object storage; state/file có thể mất khi restart/redeploy nếu không có persistent disk."),
        ("Supabase Session Pooler", "Điểm kết nối pool cho DATABASE_URL; yêu cầu TLS.", "Đích Prisma SQL qua kết nối TLS/sslmode=require."),
        ("PostgreSQL database", "Lưu users, roles, reservations, payments, wallets, parking sessions, subscriptions, incidents và dữ liệu nghiệp vụ.", "System of record."),
        ("PayOS", "Tạo/truy vấn payment link và gửi signed webhook.", "HTTPS API; webhook được backend xác minh trước khi cập nhật trạng thái."),
        ("Resend", "Gửi email HTML chứa OTP cho đăng ký/reset.", "HTTPS API/SDK; không dùng SMTP trong luồng đã xác định."),
        ("Plate Recognizer", "Nhận ảnh và trả kết quả nhận diện biển số.", "HTTPS `v1/plate-reader`; vùng VN, confidence mặc định 0.75."),
        ("UptimeRobot", "Gửi GET `/health` định kỳ.", "Chỉ kiểm tra HTTP liveness, không kiểm tra DB hoặc vendor integrations."),
        ("Build & deployment", "npm build, Prisma generate và Nest build/deploy.", "Thuộc vùng Render; repository không có CI/CD workflow committed."),
    ], widths=[1.45, 3.35, 2.05], font_size=8.1)

    doc.add_heading("4. Luồng dữ liệu và giao thức", level=1)
    doc.add_heading("4.1 Request đồng bộ từ ứng dụng", level=2)
    for item in [
        "Người dùng thao tác trên Flutter UI; Dio gửi request JSON qua HTTPS REST tới NestJS/Express `/api/*`.",
        "Backend parse/validate DTO bằng ValidationPipe. Global AtGuard kiểm tra JWT Bearer; các controller/route áp dụng role guard khi cần.",
        "Controller gọi domain service tương ứng: identity, parking/catalog, booking/subscription, billing hoặc operations.",
        "Domain service dùng Prisma để đọc/ghi PostgreSQL qua Supabase Session Pooler với TLS. Response JSON quay lại Flutter qua HTTPS.",
    ]:
        add_numbered(doc, item)

    doc.add_heading("4.2 OTP và filesystem local", level=2)
    add_table(doc, ["Luồng", "Chiều", "Giao thức / dữ liệu"], [
        ("OTP", "NestJS process ↔ Volatile local state", "Auth sử dụng `OtpStoreService` và Map trong bộ nhớ; TTL 10 phút. OTP email được gửi ra Resend qua HTTPS."),
        ("Upload", "API/controllers → `UPLOAD_DIR`", "Multipart file được ghi vào filesystem bằng Files service. `UPLOAD_DIR` mặc định là `./public/uploads`."),
        ("Static file", "Client → Nest/Express → filesystem", "Express static assets phục vụ đường dẫn `/uploads/*`; backend trả URL ảnh cho các nghiệp vụ liên quan."),
    ], widths=[1.05, 2.1, 3.7])
    add_callout(
        doc,
        "Tính chất volatile",
        "OTP Map mất khi process khởi động lại. Filesystem local trên hosting không được mặc định xem là lưu trữ bền vững; cần xác minh/cấu hình persistent disk hoặc chuyển sang object storage nếu cần giữ ảnh qua redeploy.",
        PALE_AMBER,
    )

    doc.add_heading("4.3 Thanh toán và webhook", level=2)
    add_bullet(doc, "Backend → PayOS: HTTPS API để tạo/truy vấn payment request hoặc payment link.")
    add_bullet(doc, "PayOS → backend: HTTPS POST tới `/api/payments/payos-webhook`; backend xác minh chữ ký trước khi cập nhật giao dịch trong PostgreSQL.")
    add_bullet(doc, "Webhook là callback HTTP đồng bộ tới API, không phải message queue hay event stream nội bộ.")

    doc.add_heading("4.4 OTP email và OCR", level=2)
    add_bullet(doc, "Identity/Auth → Resend: HTTPS API/SDK gửi email HTML cho OTP đăng ký hoặc reset mật khẩu.")
    add_bullet(doc, "Parking operations → Plate Recognizer: HTTPS gửi ảnh snapshot tới `v1/plate-reader`; kết quả biển số trả về backend để tiếp tục xử lý nghiệp vụ.")

    doc.add_heading("4.5 Batch và monitoring", level=2)
    add_table(doc, ["Luồng", "Tần suất / giao thức", "Ý nghĩa"], [
        ("Reservation cron", "Nest Schedule · mỗi phút", "Cron chạy trong cùng NestJS process, xử lý trạng thái reservation và cập nhật DB qua Prisma."),
        ("Health monitoring", "UptimeRobot → HTTPS GET `/health`", "Liveness HTTP từ bên ngoài; không phải health check sâu cho PostgreSQL/PayOS/Resend/OCR."),
        ("Build/deploy", "Render build/deployment → khởi chạy NestJS", "Luồng triển khai thuộc Render; không phát hiện CI/CD workflow được commit trong repository."),
    ], widths=[1.5, 2.25, 3.1])

    doc.add_heading("5. Lưu trữ dữ liệu", level=1)
    add_table(doc, ["Kho", "Loại dữ liệu", "Độ bền / kết nối"], [
        ("Supabase PostgreSQL", "Dữ liệu nghiệp vụ: tài khoản, vai trò, reservation, payment, wallet, parking session, subscription, incident và record liên quan.", "Bền vững; Prisma Client → Session Pooler → PostgreSQL qua TLS (`sslmode=require`)."),
        ("OTP Map", "OTP hoặc dữ liệu đăng ký/reset tạm thời; TTL 10 phút.", "In-memory theo process; không dùng chung giữa nhiều instance, mất khi restart."),
        ("`UPLOAD_DIR` filesystem", "Ảnh avatar, ảnh parking/QR/OCR và proof/incident theo các luồng upload.", "Filesystem local; tính bền vững phụ thuộc persistent disk và cấu hình hosting."),
    ], widths=[1.55, 3.1, 2.2])
    doc.add_paragraph("Không có bằng chứng trong codebase về Redis, cache phân tán, search index, object storage hoặc message broker/stream.")

    doc.add_heading("6. Bảo mật", level=1)
    add_table(doc, ["Kiểm soát", "Mô tả"], [
        ("Authentication", "JWT Bearer access/refresh; global AtGuard bảo vệ mặc định, route `@Public()` là ngoại lệ."),
        ("Authorization", "Role-based checks/guards trên các route nghiệp vụ được bảo vệ."),
        ("Validation", "ValidationPipe bật `whitelist` và `forbidNonWhitelisted` để kiểm soát DTO/input."),
        ("Secrets", "JWT secrets, DATABASE_URL và API keys PayOS/Resend/OCR lấy từ environment; tài liệu không ghi giá trị secrets."),
        ("Transport", "HTTPS cho client và vendor APIs; Prisma kết nối PostgreSQL qua TLS/sslmode=require."),
        ("Webhook", "PayOS webhook cần xác minh chữ ký trước khi chấp nhận cập nhật trạng thái giao dịch."),
        ("CORS", "Origin cấu hình bằng `CORS_ORIGIN`; không giả định danh sách origin cụ thể nếu chưa đọc cấu hình triển khai."),
        ("Thiết bị", "Ứng dụng lưu token trong secure storage phía mobile; không đồng nghĩa dữ liệu nghiệp vụ được lưu local."),
    ], widths=[1.35, 5.5])

    doc.add_heading("7. Vận hành, monitoring và giới hạn hạ tầng", level=1)
    add_bullet(doc, "Backend là một Render production Web Service chạy NestJS/Express; cron chạy cùng process.")
    add_bullet(doc, "UptimeRobot chỉ gọi `/health`; trong repository chưa thấy centralized logging, metrics, distributed tracing hoặc alerting tích hợp.")
    add_bullet(doc, "Không có worker riêng hoặc queue để retry/đệm các tác vụ email, OCR hay thanh toán; các integration được gọi trực tiếp qua HTTPS.")
    add_bullet(doc, "Không có Dockerfile/Compose, Kubernetes, Terraform hoặc Render blueprint được xác nhận trong source hiện tại.")
    add_bullet(doc, "Quy trình build/deploy được mô tả trong tài liệu/diagram, nhưng không có CI workflow committed; không nên mô tả như pipeline tự động đã triển khai.")
    add_bullet(doc, "Một process duy nhất cùng state in-memory có nghĩa scale nhiều instance cần xem xét OTP affinity/shared storage, cron duplication và persistence của upload.")

    doc.add_heading("8. Các thành phần không có trong runtime", level=1)
    add_table(doc, ["Hạng mục", "Trạng thái theo codebase"], [
        ("Web frontend / admin client", "Không nằm trong repository backend được mô tả; client trong phạm vi sơ đồ là Flutter mobile."),
        ("API gateway độc lập", "Không có; NestJS/Express là HTTP entry point."),
        ("Worker / queue / stream / event bus", "Không phát hiện. Cron chạy trong NestJS process."),
        ("Redis/cache, search", "Không phát hiện trong kiến trúc runtime."),
        ("Object storage", "Không phát hiện; upload hiện dùng `UPLOAD_DIR` filesystem."),
        ("WebSocket / gRPC", "Không được xác nhận trong luồng kiến trúc hiện tại; REST/HTTPS là giao thức API chính."),
        ("LLM/AI, analytics", "Không phát hiện. Plate Recognizer là dịch vụ OCR bên thứ ba, không phải LLM runtime."),
        ("Container orchestration / CI/CD", "Không có cấu hình committed được xác nhận."),
    ], widths=[2.25, 4.6])

    doc.add_heading("9. Bằng chứng trong codebase", level=1)
    add_table(doc, ["Vị trí", "Thông tin được xác nhận"], [
        ("`src/main.ts`", "Khởi chạy Nest/Express, ValidationPipe, Swagger, CORS, `UPLOAD_DIR` và static assets."),
        ("`src/app.module.ts`", "Module composition, ScheduleModule và global guard."),
        ("`src/auth/services/otp-store.service.ts`", "OTP store dùng Map trong bộ nhớ."),
        ("`src/auth/pbms/pbms-auth.service.ts`", "Đọc/ghi/xóa OTP cho đăng ký và reset; TTL."),
        ("`src/reservations/reservations.service.ts`", "Tác vụ reservation theo lịch."),
        ("`src/payments/payments.controller.ts`", "Endpoint webhook PayOS."),
        ("`src/integrations/` và `src/auth/services/mail.service.ts`", "PayOS, Plate Recognizer, file handling và Resend integrations."),
        ("`docs/huong-dan-env-va-deploy.md`, `README.md`", "Environment, Render/Supabase deployment và các lưu ý persistence."),
        ("`docs/architecture/system-architecture.json`", "Định nghĩa các node, boundary, connections và nội dung diagram."),
    ], widths=[3.0, 3.85], font_size=8.4)
    p = doc.add_paragraph()
    p.add_run("Sơ đồ tương tác và file nguồn nằm cạnh tài liệu DOCX: ").bold = True
    add_hyperlink(p, "system-architecture.html", "system-architecture.html")
    p.add_run(" / system-architecture.json")

    doc.add_heading("10. Thuật ngữ", level=1)
    add_table(doc, ["Thuật ngữ", "Ý nghĩa trong sơ đồ"], [
        ("REST / HTTPS", "Request/response HTTP mã hóa TLS giữa Flutter, backend và các dịch vụ ngoài."),
        ("Webhook", "HTTP callback từ PayOS về backend sau sự kiện thanh toán."),
        ("Batch / cron", "Tác vụ chạy theo lịch mỗi phút bên trong NestJS process."),
        ("Process-local state", "Dữ liệu nằm trong RAM hoặc filesystem của process/instance hiện tại, không phải kho dùng chung."),
        ("System of record", "PostgreSQL là nguồn dữ liệu nghiệp vụ lưu bền vững chính."),
        ("Trust boundary", "Ranh giới ứng dụng/secrets; không phải một service riêng hay ranh giới cloud độc lập."),
    ], widths=[1.65, 5.2])

    doc.add_paragraph()
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = end.add_run("— Hết —")
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(MUTED)
    run.font.size = Pt(9)

    doc.core_properties.title = "PRM393 Parking Management — System Architecture"
    doc.core_properties.subject = "Vietnamese architecture document based on the system architecture diagram and repository evidence"
    doc.core_properties.author = "PRM393 Project"
    doc.core_properties.keywords = "Flutter, NestJS, Render, Supabase, PostgreSQL, PayOS, Resend, Plate Recognizer"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
